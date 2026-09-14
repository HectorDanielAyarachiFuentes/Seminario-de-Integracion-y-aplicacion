import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document("Entregas/Actividad 1/TP1_Ayarachi_Fuentes.docx")

# Find pb_idx
pb_idx = None
p_logo = None
for i, p in enumerate(doc.paragraphs):
    if len(p._element.xpath('.//w:drawing')) > 0 and p_logo is None:
        p_logo = p
    if 'type="page"' in p._p.xml or p.text.strip() == "Índice de Contenidos":
        pb_idx = i if 'type="page"' in p._p.xml else i - 1
        break

print("pb_idx:", pb_idx)
print("Paragraphs before pb:")
for p in list(doc.paragraphs[:pb_idx]):
    txt = p.text.strip()
    if not txt and len(p._element.xpath('.//w:drawing')) == 0:
        print("Removing empty paragraph:", repr(p.text))
        p._p.getparent().remove(p._p)
    else:
        print("Keeping paragraph:", repr(txt[:30]))
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)

doc.save("Entregas/Actividad 1/test_fixed.docx")
print("Saved test_fixed.docx successfully")
