import pymupdf
import docx

print("==================================================")
print("VERIFICACIÓN PDF: TP1_Ayarachi_Fuentes.pdf")
print("==================================================")
doc_pdf = pymupdf.open("Entregas/Actividad 1/TP1_Ayarachi_Fuentes.pdf")
print(f"Total de páginas: {len(doc_pdf)}")
for i, page in enumerate(doc_pdf):
    footers = [b[4].strip() for b in page.get_text("blocks") if b[1] > page.rect.height * 0.85]
    first_line = page.get_text().strip().split("\n")[0] if page.get_text().strip() else ""
    print(f"  Pág. {i}: '{first_line[:40]}' | Pie: {footers}")

print("\n--- Texto del Índice de Contenidos (Pág. 1) ---")
print(doc_pdf[1].get_text())

print("\n==================================================")
print("VERIFICACIÓN WORD: TP1_Ayarachi_Fuentes.docx")
print("==================================================")
doc_word = docx.Document("Entregas/Actividad 1/TP1_Ayarachi_Fuentes.docx")
print(f"Total párrafos: {len(doc_word.paragraphs)}")
print(f"Total tablas: {len(doc_word.tables)}")
for i, table in enumerate(doc_word.tables):
    print(f"  Tabla {i+1}: {len(table.rows)} filas x {len(table.columns)} cols")
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    borders = table._tbl.tblPr.xpath(".//w:tblBorders")
    if borders:
        print(f"    Bordes: insideH={borders[0].xpath('.//w:insideH/@w:val', namespaces=ns)}, insideV={borders[0].xpath('.//w:insideV/@w:val', namespaces=ns)}")

print("\nPárrafos bajo Referencias Bibliográficas en Word:")
is_ref = False
for p in doc_word.paragraphs:
    if "Referencias Bibliográficas" in p.text:
        is_ref = True
        print(f"  Encabezado: '{p.text}'")
        continue
    if is_ref and p.text.strip():
        print(f"  - '{p.text[:60]}...' (left={p.paragraph_format.left_indent}, first_line={p.paragraph_format.first_line_indent})")

