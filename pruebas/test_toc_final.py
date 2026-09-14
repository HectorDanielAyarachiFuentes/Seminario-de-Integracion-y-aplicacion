import docx
import shutil
import pymupdf
from pathlib import Path
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

src_docx = Path('Entregas/Actividad 1/TP1_Ayarachi_Fuentes.docx')
pdf_path = Path('Entregas/Actividad 1/TP1_Ayarachi_Fuentes.pdf')
dst_docx = Path('pruebas/test_toc_final.docx')
shutil.copy2(src_docx, dst_docx)

doc_word = docx.Document(str(dst_docx))

# 1. Extraer ítems desde el PDF (página de índice de Typst)
items_toc = []
if pdf_path.exists():
    try:
        doc_pdf = pymupdf.open(str(pdf_path))
        if len(doc_pdf) > 1:
            txt_toc = doc_pdf[1].get_text()
            for line in txt_toc.split('\n'):
                line = line.strip()
                if not line or line.startswith('Índice') or line == '1':
                    continue
                parts = line.split(' .')
                if len(parts) >= 2:
                    title = parts[0].strip()
                    pag = parts[-1].replace('.', '').replace('\u2060', '').strip()
                    if pag.isdigit():
                        level = 2 if title == 'Objetivos' or title.startswith('1.1') else 1
                        items_toc.append((title, level, pag))
    except Exception as e:
        print("Error leyendo PDF:", e)

# 2. Obtener los textos de títulos de Word (con numeración)
headings_word = []
for p in doc_word.paragraphs:
    txt = p.text.strip()
    if txt == "Índice de Contenidos":
        continue
    if p.style.name.startswith(('Heading', 'Encabezado')) or txt.startswith(('1.', '2.', '3.', '4.', '5.')):
        headings_word.append(txt)

# 3. Cruzar títulos para que en Word tengan la numeración de sección
final_items = []
for title_pdf, level, pag in items_toc:
    matched = title_pdf
    for hw in headings_word:
        if title_pdf in hw or hw.endswith(title_pdf):
            matched = hw
            break
    final_items.append((matched, level, pag))

print("Final TOC items:")
for it in final_items:
    print(it)

# 4. Localizar 'Índice de Contenidos' y limpiar cualquier TOC anterior
p_toc = None
for p in doc_word.paragraphs:
    if p.text.strip() == "Índice de Contenidos":
        p_toc = p
        break

if p_toc:
    # Eliminar párrafos siguientes si eran campos TOC viejos
    next_p = p_toc._p.getnext()
    while next_p is not None:
        xml_s = next_p.xml
        if 'TOC' in xml_s or 'fldChar' in xml_s:
            parent = next_p.getparent()
            temp = next_p.getnext()
            parent.remove(next_p)
            next_p = temp
        else:
            break

    # Insertar el campo TOC completo con entradas pre-renderizadas
    tab_pos = "9026"
    curr = p_toc._p

    begin_xml = (
        f'<w:p {nsdecls("w")}>'
        f'<w:pPr><w:spacing w:after="120"/></w:pPr>'
        f'<w:r>'
        f'<w:fldChar w:fldCharType="begin"/>'
        f'<w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>'
        f'<w:fldChar w:fldCharType="separate"/>'
        f'</w:r>'
        f'</w:p>'
    )
    p_begin = parse_xml(begin_xml)
    curr.addnext(p_begin)
    curr = p_begin

    for texto, nivel, num_pag in final_items:
        style_name = f"TOC{nivel}"
        left_indent = "280" if nivel == 2 else "0"
        item_xml = (
            f'<w:p {nsdecls("w")}>'
            f'<w:pPr>'
            f'<w:pStyle w:val="{style_name}"/>'
            f'<w:tabs>'
            f'<w:tab w:val="right" w:leader="dot" w:pos="{tab_pos}"/>'
            f'</w:tabs>'
            f'<w:ind w:left="{left_indent}"/>'
            f'<w:spacing w:after="60"/>'
            f'</w:pPr>'
            f'<w:r>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Segoe UI" w:hAnsi="Segoe UI"/>'
            f'<w:sz w:val="20"/>'
            f'</w:rPr>'
            f'<w:t>{texto}</w:t>'
            f'</w:r>'
            f'<w:r>'
            f'<w:tab/>'
            f'</w:r>'
            f'<w:r>'
            f'<w:rPr>'
            f'<w:rFonts w:ascii="Segoe UI" w:hAnsi="Segoe UI"/>'
            f'<w:sz w:val="20"/>'
            f'</w:rPr>'
            f'<w:t>{num_pag}</w:t>'
            f'</w:r>'
            f'</w:p>'
        )
        p_item = parse_xml(item_xml)
        curr.addnext(p_item)
        curr = p_item

    end_xml = (
        f'<w:p {nsdecls("w")}>'
        f'<w:r>'
        f'<w:fldChar w:fldCharType="end"/>'
        f'</w:r>'
        f'</w:p>'
    )
    p_end = parse_xml(end_xml)
    curr.addnext(p_end)

# 5. Activar updateFields en settings.xml
try:
    settings = doc_word.settings.element
    update_fields = parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>')
    settings.append(update_fields)
    print("updateFields activado")
except Exception as e:
    print("Error updateFields:", e)

doc_word.save(str(dst_docx))
print("Archivo final guardado exitosamente.")
